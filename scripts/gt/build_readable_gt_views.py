import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def load_jsonl(path):
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def as_text(value):
    if value is None:
        return ""
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def safe_join(values):
    if not values:
        return ""
    return "；".join(str(value) for value in values if value is not None)


def amount(record, name):
    calculation = record.get("calculation_gt") or {}
    value = calculation.get(name)
    return "" if value is None else value


def case_total(record):
    return amount(record, "total_reimbursable")


def route_text(record):
    rows = []
    for item in record.get("train_ticket_summary") or []:
        rows.append(
            "{person}: {date} {from_city}->{to_city} {train_no} {amount}".format(
                person=item.get("person", ""),
                date=item.get("date", ""),
                from_city=item.get("from_city", ""),
                to_city=item.get("to_city", ""),
                train_no=item.get("train_no", ""),
                amount=item.get("amount", ""),
            )
        )
    return "；".join(rows)


def hotel_text(record):
    rows = []
    for item in record.get("hotel_invoice_summary") or []:
        rows.append(
            "{city}: {checkin}->{checkout}, {nights}晚, {amount}, {guests}".format(
                city=item.get("city", ""),
                checkin=item.get("checkin", ""),
                checkout=item.get("checkout", ""),
                nights=item.get("nights", ""),
                amount=item.get("amount", ""),
                guests=safe_join(item.get("guests") or []),
            )
        )
    return "；".join(rows)


def first_slots(record, limit=20):
    filled = record.get("filled_template_gt")
    if not filled:
        return []
    slots = []
    for slot in filled.get("traveler_name_cells") or []:
        slots.append(slot)
    for row in filled.get("transport_detail_rows") or []:
        for name in ("departure_date", "departure_city", "arrival_date", "arrival_city", "transport", "amount"):
            if name in row:
                slots.append(row[name])
    for group in (filled.get("fixed_expense_rows") or {}).values():
        for name in ("person_count", "day_count", "amount"):
            if name in group:
                slots.append(group[name])
    for name in ("total", "total_explanation"):
        if name in filled:
            slots.append(filled[name])
    return slots[:limit]


def write_pretty_json(record, out_dir):
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "case.pretty.json").write_text(
        json.dumps(record, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def write_markdown(record, out_dir):
    out_dir.mkdir(parents=True, exist_ok=True)
    lines = [
        "# {case_id}".format(case_id=record["case_id"]),
        "",
        "## 基本信息",
        "",
        "| 字段 | 值 |",
        "|---|---|",
        "| subset | {0} |".format(record.get("subset", "")),
        "| subset_name | {0} |".format(record.get("subset_name", "")),
        "| scenario_family | {0} |".format(record.get("scenario_family", "")),
        "| scenario_type | {0} |".format(record.get("scenario_type", "")),
        "| gold_action | {0} |".format(record.get("gold_action", "")),
        "| action_code | {0} |".format(record.get("action_code", "")),
        "| person_count | {0} |".format(record.get("person_count", "")),
        "| trip_start | {0} |".format(record.get("trip_start", "")),
        "| trip_end | {0} |".format(record.get("trip_end", "")),
        "| travelers | {0} |".format(safe_join(record.get("resolved_travelers_gt") or record.get("evidence_persons") or [])),
        "",
        "## 证据摘要",
        "",
        "### 车票",
        "",
    ]
    train = record.get("train_ticket_summary") or []
    if train:
        lines += ["| # | 人员 | 日期 | 出发 | 到达 | 车次 | 席别 | 金额 | 状态 |", "|---:|---|---|---|---|---|---|---:|---|"]
        for item in train:
            lines.append(
                "| {leg_no} | {person} | {date} | {from_city} | {to_city} | {train_no} | {seat} | {amount} | {status} |".format(
                    leg_no=item.get("leg_no", ""),
                    person=item.get("person", ""),
                    date=item.get("date", ""),
                    from_city=item.get("from_city", ""),
                    to_city=item.get("to_city", ""),
                    train_no=item.get("train_no", ""),
                    seat=item.get("seat", ""),
                    amount=item.get("amount", ""),
                    status=item.get("status", ""),
                )
            )
    else:
        lines.append("无")
    lines += ["", "### 酒店", ""]
    hotels = record.get("hotel_invoice_summary") or []
    if hotels:
        lines += ["| 城市 | 入住 | 离店 | 晚数 | 酒店 | 金额 | 入住人 | 状态 |", "|---|---|---|---:|---|---:|---|---|"]
        for item in hotels:
            lines.append(
                "| {city} | {checkin} | {checkout} | {nights} | {hotel_name} | {amount} | {guests} | {status} |".format(
                    city=item.get("city", ""),
                    checkin=item.get("checkin", ""),
                    checkout=item.get("checkout", ""),
                    nights=item.get("nights", ""),
                    hotel_name=item.get("hotel_name", ""),
                    amount=item.get("amount", ""),
                    guests=safe_join(item.get("guests") or []),
                    status=item.get("status", ""),
                )
            )
    else:
        lines.append("无")

    lines += [
        "",
        "## 金额 GT",
        "",
        "| 分项 | 金额 |",
        "|---|---:|",
        "| 交通费 | {0} |".format(amount(record, "transport_reimbursable")),
        "| 住宿费 | {0} |".format(amount(record, "hotel_reimbursable")),
        "| 伙食补助 | {0} |".format(amount(record, "meal_allowance")),
        "| 市内交通补助 | {0} |".format(amount(record, "local_transport_allowance")),
        "| 总计 | {0} |".format(amount(record, "total_reimbursable")),
    ]

    if record.get("reject_reason"):
        lines += ["", "## Reject 原因", "", record["reject_reason"]]

    if record.get("defer_reason") or record.get("defer_handoff_gt"):
        lines += ["", "## Defer 人工交接", "", "### 阻断问题", "", as_text(record.get("defer_reason"))]
        handoff = record.get("defer_handoff_gt") or {}
        if handoff:
            lines += [
                "",
                "### 已知事实",
                "",
                "```json",
                json.dumps(handoff.get("known_facts", {}), ensure_ascii=False, indent=2),
                "```",
                "",
                "### 已知金额",
                "",
                "```json",
                json.dumps(handoff.get("known_amounts", {}), ensure_ascii=False, indent=2),
                "```",
                "",
                "### 未决字段",
                "",
                safe_join(handoff.get("unresolved_fields") or []),
                "",
                "### 所需材料",
                "",
            ]
            lines += ["- {0}".format(item) for item in handoff.get("required_information") or []]

    if record.get("filled_template_gt"):
        lines += [
            "",
            "## Process 关键槽位预览",
            "",
            "| cell | slot_name | slot_value | evidence_count | trace_count | calculation |",
            "|---|---|---|---:|---:|---|",
        ]
        for slot in first_slots(record):
            lines.append(
                "| {cell} | {slot_name} | {slot_value} | {evidence_count} | {trace_count} | {calculation} |".format(
                    cell=slot.get("cell", ""),
                    slot_name=slot.get("slot_name", ""),
                    slot_value=as_text(slot.get("slot_value")),
                    evidence_count=len(slot.get("evidence_sources") or []),
                    trace_count=len(slot.get("travel_standard_reasoning_trace") or []),
                    calculation=as_text(slot.get("calculation")),
                )
            )

    lines += [
        "",
        "## 完整机器字段",
        "",
        "完整 JSON 见同目录 `case.pretty.json`。",
        "",
    ]
    (out_dir / "case.summary.md").write_text("\n".join(lines), encoding="utf-8")


def write_index(records, out_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "readable_index"
    headers = [
        "case_id",
        "subset",
        "subset_name",
        "gold_action",
        "action_code",
        "scenario_family",
        "scenario_type",
        "person_count",
        "travelers",
        "trip_start",
        "trip_end",
        "route_summary",
        "hotel_summary",
        "transport_reimbursable",
        "hotel_reimbursable",
        "meal_allowance",
        "local_transport_allowance",
        "total_reimbursable",
        "reject_reason",
        "defer_reason",
        "summary_md",
        "pretty_json",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for record in records:
        rel_dir = Path("readable_cases") / record["subset"] / record["case_id"]
        ws.append(
            [
                record.get("case_id"),
                record.get("subset"),
                record.get("subset_name"),
                record.get("gold_action"),
                record.get("action_code"),
                record.get("scenario_family"),
                record.get("scenario_type"),
                record.get("person_count"),
                safe_join(record.get("resolved_travelers_gt") or record.get("evidence_persons") or []),
                record.get("trip_start"),
                record.get("trip_end"),
                route_text(record),
                hotel_text(record),
                amount(record, "transport_reimbursable"),
                amount(record, "hotel_reimbursable"),
                amount(record, "meal_allowance"),
                amount(record, "local_transport_allowance"),
                amount(record, "total_reimbursable"),
                record.get("reject_reason"),
                record.get("defer_reason"),
                str(rel_dir / "case.summary.md"),
                str(rel_dir / "case.pretty.json"),
            ]
        )

    widths = {
        "A": 14,
        "B": 8,
        "C": 24,
        "D": 12,
        "E": 28,
        "F": 24,
        "G": 32,
        "H": 12,
        "I": 24,
        "J": 12,
        "K": 12,
        "L": 60,
        "M": 60,
        "S": 60,
        "T": 60,
        "U": 42,
        "V": 42,
    }
    for column in range(1, len(headers) + 1):
        letter = get_column_letter(column)
        ws.column_dimensions[letter].width = widths.get(letter, 16)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main():
    script_path = Path(__file__).resolve()
    repo_root = script_path.parents[2]
    gt_path = repo_root / "gt" / "cases.jsonl"
    readable_root = repo_root / "gt" / "readable_cases"
    index_path = repo_root / "gt" / "case_index.xlsx"
    summary_path = repo_root / "gt" / "readable_generation_summary.json"

    records = load_jsonl(gt_path)
    for record in records:
        case_dir = readable_root / record["subset"] / record["case_id"]
        write_pretty_json(record, case_dir)
        write_markdown(record, case_dir)
    write_index(records, index_path)

    payload = {
        "source_gt": str(gt_path.relative_to(repo_root)),
        "records": len(records),
        "actions": dict(sorted(Counter(record["gold_action"] for record in records).items())),
        "subsets": dict(sorted(Counter(record["subset"] for record in records).items())),
        "readable_cases_dir": str(readable_root.relative_to(repo_root)),
        "index_xlsx": str(index_path.relative_to(repo_root)),
        "generated_files": len(list(readable_root.rglob("*.json"))) + len(list(readable_root.rglob("*.md"))) + 2,
    }
    summary_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(payload, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
