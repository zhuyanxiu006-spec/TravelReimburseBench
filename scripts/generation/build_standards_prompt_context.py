#!/usr/bin/env python3
"""Build the standards text block inserted into action prompts.

The source standards remain in the standards/ directory. This script extracts
policy text, hotel-limit tables, and reimbursement-template slots into one text
file that runners can pass to the prompt placeholder named {policy_text}.
"""

import argparse
from pathlib import Path
from typing import Iterable, List

from docx import Document
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_POLICY = ROOT / "standards" / "reimbursement_policy.docx"
DEFAULT_HOTEL_LIMITS = ROOT / "standards" / "hotel_limits.xlsx"
DEFAULT_TEMPLATE = ROOT / "standards" / "reimbursement_form_template.xlsx"
DEFAULT_OUTPUT = ROOT / "standards" / "standards_context.txt"


def clean_cell(value):
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    return " ".join(part.strip() for part in text.split("\n") if part.strip())


def non_empty(values: Iterable[str]) -> List[str]:
    return [value for value in values if value]


def extract_docx(path: Path) -> str:
    doc = Document(str(path))
    lines = []
    for paragraph in doc.paragraphs:
        text = clean_cell(paragraph.text)
        if text:
            lines.append(text)

    for table_index, table in enumerate(doc.tables, 1):
        lines.append("")
        lines.append("【政策表格 %s】" % table_index)
        for row in table.rows:
            cells = non_empty(clean_cell(cell.text) for cell in row.cells)
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def extract_xlsx(path: Path) -> str:
    wb = load_workbook(str(path), data_only=True)
    sections = []
    for ws in wb.worksheets:
        rows = []
        last_values = {}
        for row in ws.iter_rows(values_only=True):
            raw_values = [clean_cell(value) for value in row]
            if not any(raw_values):
                continue

            values = []
            for index, value in enumerate(raw_values):
                if value:
                    last_values[index] = value
                    values.append(value)
                else:
                    # Carry merged-like leading identifiers forward only when
                    # later cells contain a sub-row. This preserves hotel-limit
                    # season rows without filling every blank cell noisily.
                    values.append(last_values.get(index, "") if any(raw_values[index + 1 :]) else "")

            while values and not values[-1]:
                values.pop()
            if values:
                rows.append(" | ".join(values))
        sections.append("【工作表：%s】\n%s" % (ws.title, "\n".join(rows)))
    return "\n\n".join(sections)


def build_context(policy: Path, hotel_limits: Path, template: Path) -> str:
    parts = [
        "# TravelReimburseBench standards context",
        "",
        "说明：以下内容由 standards 目录中的正式文件抽取生成，用于统一放入模型 prompt。模型只能依据这些标准文件和当前 case 票据材料判断，不得使用外部规则。",
        "",
        "## 1. 报销制度文件：standards/reimbursement_policy.docx",
        extract_docx(policy),
        "",
        "## 2. 酒店住宿标准：standards/hotel_limits.xlsx",
        extract_xlsx(hotel_limits),
        "",
        "## 3. 报销模板说明：standards/reimbursement_form_template.xlsx",
        "模板用于评估 process case 的字段填写。模型需要根据车票、住宿发票和标准文件填写出差人、交通明细、住宿费、伙食补助、市内交通补助和总计金额等槽位。",
        extract_xlsx(template),
        "",
    ]
    return "\n".join(parts)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--policy", type=Path, default=DEFAULT_POLICY)
    parser.add_argument("--hotel-limits", type=Path, default=DEFAULT_HOTEL_LIMITS)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    for path in [args.policy, args.hotel_limits, args.template]:
        if not path.exists():
            raise FileNotFoundError(path)

    text = build_context(args.policy, args.hotel_limits, args.template)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(text, encoding="utf-8", newline="\n")
    print("wrote=%s" % args.output)
    print("chars=%s" % len(text))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
